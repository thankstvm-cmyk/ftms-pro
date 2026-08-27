import html
import math
import os
import sqlite3
import tkinter as tk
from datetime import datetime
from tkinter import filedialog, messagebox, ttk

from PIL import Image, ImageTk
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A3, A4, B4, B5, LETTER, legal, landscape, portrait
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.platypus import Image as ReportLabImage
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from config import DATABASE_PATH
from font_manager import FontManager


class ConvertPdf:
    """PDF settings, preview, and export window for an existing report result."""

    PAGE_SIZES = {"A4": A4, "A3": A3, "LETTER": LETTER, "Legal": legal, "B4": B4, "B5": B5}

    def __init__(self, parent, controller, report_title="Vehicle PDF Report Page",
                 report_type="vehicle", columns=None, data=None,
                 table_name="vehicles", field_mapping=None):
        self.parent = parent
        self.controller = controller
        self.report_title = report_title
        self.report_type = report_type
        self.columns = list(columns or [])
        self.data = list(data or [])
        self.table_name = table_name
        self.field_mapping = field_mapping or {}
        self.pdf_field_vars = {}
        self.logo_img = None

    def get_company_details(self):
        try:
            with sqlite3.connect(DATABASE_PATH) as conn:
                row = conn.execute("SELECT name, logo FROM company LIMIT 1").fetchone()
            return row if row else ("", "")
        except sqlite3.Error:
            return "", ""

    def open_pdf_settings(self):
        self.pdf_win = tk.Toplevel(self.parent)
        self.pdf_win.title(self.report_title)
        self.pdf_win.state("zoomed")

        self.top_frame = tk.Frame(self.pdf_win)
        self.top_frame.pack(fill="both", expand=True)
        self.bottom_frame = tk.Frame(self.pdf_win, bg="#d8e6ec", height=50)
        self.bottom_frame.pack(side="bottom", fill="x")
        self.bottom_frame.pack_propagate(False)
        container = tk.Frame(self.top_frame)
        container.pack(fill="both", expand=True)

        for column, weight in enumerate((2, 3, 3, 2)):
            container.grid_columnconfigure(column, weight=weight)
        container.grid_rowconfigure(0, weight=1)
        self.left_frame1 = ttk.LabelFrame(container, text="SELECT FIELDS")
        self.left_frame2 = ttk.LabelFrame(container, text="PAGE SETUP")
        self.right_frame1 = ttk.LabelFrame(container, text="HEADING, TITLE & BODY")
        self.right_frame2 = ttk.LabelFrame(container, text="MORE SETTINGS")
        for column, frame in enumerate((self.left_frame1, self.left_frame2, self.right_frame1, self.right_frame2)):
            frame.grid(row=0, column=column, sticky="nsew", padx=2, pady=2)
            frame.grid_propagate(False)

        self._create_field_controls()
        self._create_page_controls()
        self._create_typography_controls()
        self._create_summary_controls()

        button_frame = tk.Frame(self.bottom_frame)
        button_frame.pack(pady=3)
        tk.Button(button_frame, text="PRINT PREVIEW", font=("Segoe UI", 10, "bold"), bg="#1f8b4c", fg="white",
                  activebackground="#146c37", width=20, command=self.show_preview).grid(row=0, column=0, padx=10)
        tk.Button(button_frame, text="PRINT", font=("Segoe UI", 10, "bold"), bg="#1f8b4c", fg="white",
                  activebackground="#146c37", width=20, command=self.print_pdf).grid(row=0, column=1, padx=10)
        tk.Button(button_frame, text="SAVE AS PDF", font=("Segoe UI", 10, "bold"), bg="#1f8b4c", fg="white",
                  activebackground="#146c37", width=20, command=self.generate_pdf).grid(row=0, column=2, padx=10)
        tk.Button(button_frame, text="SAVE AS EXCEL", font=("Segoe UI", 10, "bold"), bg="#1f8b4c", fg="white",
                  activebackground="#146c37", width=20, command=self.generate_excel).grid(row=0, column=3, padx=10)
        tk.Button(button_frame, text="CLOSE/CANCEL", font=("Segoe UI", 10, "bold"), bg="#1f8b4c", fg="white",
                  activebackground="#146c37", width=20, command=self.pdf_win_close).grid(row=0, column=4, padx=10)

        self._watch_settings()
        self.update_field_count()
        self.show_preview()

    def _create_field_controls(self):
        for index, column in enumerate(self.columns):
            var = tk.BooleanVar(value=False)
            self.pdf_field_vars[column] = var
            tk.Checkbutton(self.left_frame1, text=column, variable=var,
                           command=self.update_field_count).grid(row=index, column=0, sticky="w", padx=10)
        button_row = len(self.columns) + 1
        buttons = tk.Frame(self.left_frame1)
        buttons.grid(row=button_row, column=0, pady=4)
        tk.Button(buttons, text="SELECT ALL", font=("Segoe UI", 9, "bold"), width=15,
                  command=self.select_all_fields).pack(side="left", padx=5)
        tk.Button(buttons, text="CLEAR ALL", font=("Segoe UI", 9, "bold"), width=15,
                  command=self.clear_all_fields).pack(side="left", padx=5)
        self.lbl_field_count = tk.Label(self.left_frame1, font=("Segoe UI", 9, "bold"), justify="left")
        self.lbl_field_count.grid(row=button_row + 1, column=0, sticky="w", padx=10, pady=(2, 6))

    def _create_page_controls(self):
        paper = ttk.LabelFrame(self.left_frame2, text="Paper Settings")
        paper.grid(row=0, column=0, padx=10, pady=10, sticky="ew")
        self.orientation_var = tk.StringVar(value="Landscape")
        self.page_size_var = tk.StringVar(value="A4")
        for row, label, variable, values in (
            (0, "Page Orientation", self.orientation_var, ("Landscape", "Portrait")),
            (1, "Page Size", self.page_size_var, tuple(self.PAGE_SIZES)),
        ):
            tk.Label(paper, text=label, width=20, anchor="w").grid(row=row, column=0, padx=10, pady=5, sticky="w")
            combo = ttk.Combobox(paper, textvariable=variable, values=values, state="readonly", width=20)
            combo.grid(row=row, column=1, padx=10, pady=5, sticky="w")
            combo.bind("<<ComboboxSelected>>", lambda _event: self.update_field_count())

        margins = ttk.LabelFrame(self.left_frame2, text="Margin Settings")
        margins.grid(row=1, column=0, padx=10, pady=10, sticky="ew")
        self.top_margin_var = tk.IntVar(value=10)
        self.bottom_margin_var = tk.IntVar(value=10)
        self.left_margin_var = tk.IntVar(value=10)
        self.right_margin_var = tk.IntVar(value=10)
        for row, label, variable in (
            (0, "Top Margin", self.top_margin_var), (1, "Bottom Margin", self.bottom_margin_var),
            (2, "Left Margin", self.left_margin_var), (3, "Right Margin", self.right_margin_var),
        ):
            tk.Label(margins, text=label, width=15, anchor="w").grid(row=row, column=0, padx=10, pady=2, sticky="w")
            ttk.Spinbox(margins, from_=0, to=50, textvariable=variable, width=10).grid(row=row, column=1, padx=10, pady=2, sticky="w")

        self.preview_frame = tk.LabelFrame(self.left_frame2, text="Print Preview", bg="white")
        self.preview_frame.grid(row=2, column=0, padx=5, pady=5, sticky="n")
        self.preview_canvas = tk.Canvas(self.preview_frame, width=320, height=253, bg="skyblue")
        self.preview_canvas.pack(padx=2, pady=2)
        self.preview_info_label = tk.Label(self.preview_frame, text="", font=("Arial", 8, "bold"),
                                           justify="left", wraplength=315, bg="#f3f69b")
        self.preview_info_label.pack(fill="x", padx=2, pady=(0, 2))

    def _create_typography_controls(self):
        fonts = self._available_fonts()
        self.company_heading_var = tk.BooleanVar(value=False)
        self.company_logo_var = tk.BooleanVar(value=False)
        self.show_date_var = tk.BooleanVar(value=True)
        self.show_page_var = tk.BooleanVar(value=True)
        self.company_title_var = tk.StringVar()
        self.font_comphead_var = tk.StringVar(value="Helvetica")
        self.font_compreport_var = tk.StringVar(value="Helvetica")
        self.font_body_var = tk.StringVar(value="Helvetica")
        self.font_size_comphead_var = tk.IntVar(value=13)
        self.font_size_compreport_var = tk.IntVar(value=12)
        self.font_size_body_var = tk.IntVar(value=9)

        company = ttk.LabelFrame(self.right_frame1, text="Company Name - Settings")
        company.grid(row=0, column=0, padx=10, pady=10, sticky="ew")
        tk.Checkbutton(company, text="Include Company Name", variable=self.company_heading_var).grid(row=0, column=0, padx=10, pady=3, sticky="w")
        tk.Checkbutton(company, text="Include Company Logo", variable=self.company_logo_var,
                       command=self.toggle_company_logo).grid(row=1, column=0, padx=10, pady=3, sticky="w")
        self._add_font_controls(company, 2, "Font", self.font_comphead_var, self.font_size_comphead_var, fonts)

        title = ttk.LabelFrame(self.right_frame1, text="Report Title - Settings")
        title.grid(row=1, column=0, padx=10, pady=10, sticky="ew")
        tk.Label(title, text="Report Title", width=20, anchor="w").grid(row=0, column=0, padx=10, pady=5, sticky="w")
        tk.Entry(title, textvariable=self.company_title_var, width=30).grid(row=0, column=1, padx=10, pady=5, sticky="w")
        self._add_font_controls(title, 1, "Font", self.font_compreport_var, self.font_size_compreport_var, fonts)

        body = ttk.LabelFrame(self.right_frame1, text="Report Body Settings")
        body.grid(row=2, column=0, padx=10, pady=10, sticky="ew")
        self._add_font_controls(body, 0, "Font", self.font_body_var, self.font_size_body_var, fonts)

        logo = ttk.LabelFrame(self.right_frame1, text="Company Logo")
        logo.grid(row=3, column=0, padx=5, pady=5)
        self.logo_canvas = tk.Canvas(logo, bg="white", relief="sunken", bd=1, width=300, height=145)
        self.logo_canvas.pack(padx=10, pady=10)
        self.clear_logo_preview()

    def _add_font_controls(self, parent, row, label, font_var, size_var, fonts):
        tk.Label(parent, text=label, width=20, anchor="w").grid(row=row, column=0, padx=10, pady=5, sticky="w")
        ttk.Combobox(parent, textvariable=font_var, values=fonts, state="normal", width=20).grid(row=row, column=1, padx=10, pady=5, sticky="ew")
        tk.Label(parent, text="Font Size", width=20, anchor="w").grid(row=row + 1, column=0, padx=10, pady=5, sticky="w")
        ttk.Spinbox(parent, from_=6, to=20, textvariable=size_var, width=20).grid(row=row + 1, column=1, padx=10, pady=5, sticky="ew")

    def _create_summary_controls(self):
        options = ttk.LabelFrame(self.right_frame2, text="Include Date & P-Number in the Report")
        options.grid(row=0, column=0, padx=10, pady=10, sticky="ew")
        tk.Checkbutton(options, text="Show Date & Time", width=20, variable=self.show_date_var).grid(row=0, column=0, padx=10, pady=5, sticky="w")
        tk.Checkbutton(options, text="Show Page Number", width=20, variable=self.show_page_var).grid(row=1, column=0, padx=10, pady=5, sticky="w")
        summary = ttk.LabelFrame(self.right_frame2, text="Report Statistics")
        summary.grid(row=1, column=0, padx=10, pady=10, sticky="ew")
        self.summary_labels = {}
        for row, (label, key) in enumerate((
            ("Report Name :", "name"), ("Records :", "records"), ("Selected Fields :", "fields"),
            ("Page Size :", "page_size"), ("Orientation :", "orientation"), ("Capacity/Page :", "capacity"),
            ("Overflow :", "overflow"), ("Estimated Pages :", "pages"), ("Status :", "status"),
        )):
            ttk.Label(summary, text=label, anchor="w", width=18).grid(row=row, column=0, padx=5, pady=3, sticky="w")
            value = ttk.Label(summary, text="", anchor="w")
            value.grid(row=row, column=1, padx=5, pady=3, sticky="w")
            self.summary_labels[key] = value
        guidance = tk.LabelFrame(self.right_frame2, text="FTMS FLEETPRO REPORT GUIDANCE", font=("Arial", 10, "bold"), padx=8, pady=8)
        guidance.grid(row=2, column=0, padx=10, pady=10, sticky="nsew")
        self.guidance_label = tk.Label(guidance, justify="left", anchor="nw", wraplength=250, font=("Segoe UI", 8))
        self.guidance_label.pack(fill="both", expand=True)

    def _watch_settings(self):
        for variable in (self.orientation_var, self.page_size_var, self.top_margin_var, self.bottom_margin_var,
                         self.left_margin_var, self.right_margin_var, self.company_heading_var, self.company_logo_var,
                         self.company_title_var, self.font_comphead_var, self.font_compreport_var, self.font_body_var,
                         self.font_size_comphead_var, self.font_size_compreport_var, self.font_size_body_var,
                         self.show_date_var, self.show_page_var):
            variable.trace_add("write", lambda *_args: self.update_field_count())

    def _available_fonts(self):
        return sorted(set(["Helvetica", "Times-Roman", "Courier"] + FontManager.get_font_names()))

    def _selected_columns(self):
        return [column for column, variable in self.pdf_field_vars.items() if variable.get()]

    def _selected_report_data(self):
        """Return the selected headings and matching cells from the visible report rows."""
        selected = self._selected_columns()
        indexes = [self.columns.index(column) for column in selected]
        rows = [
            [row[index] if index < len(row) and row[index] is not None else "" for index in indexes]
            for row in self.data
        ]
        return selected, rows

    def _build_report_context(self):
        """Build the single report definition consumed by both output formats."""
        selected, report_rows = self._selected_report_data()
        company_name, logo_path = self.get_company_details()
        return {
            "selected": selected,
            "rows": report_rows,
            "metrics": self._layout_metrics(),
            "title": self.company_title_var.get().strip() or self.report_title,
            "company_name": company_name,
            "logo_path": logo_path,
            "generated_at": datetime.now(),
            "company_font": self.font_comphead_var.get(),
            "title_font": self.font_compreport_var.get(),
            "body_font": self.font_body_var.get(),
        }

    def _layout_metrics(self):
        selected = self._selected_columns()
        body_size = max(6, int(self.font_size_body_var.get()))
        page_size = self.PAGE_SIZES[self.page_size_var.get()]
        page_size = landscape(page_size) if self.orientation_var.get() == "Landscape" else portrait(page_size)
        width, height = page_size
        top = int(self.top_margin_var.get()) * mm
        bottom = int(self.bottom_margin_var.get()) * mm
        printable_height = max(1, height - top - bottom)
        header_height = 22 + (int(self.font_size_comphead_var.get()) + 8 if self.company_heading_var.get() else 0)
        header_height += 26 if self.company_logo_var.get() else 0
        row_height = max(12, body_size * 1.75)
        capacity = max(1, int((printable_height - header_height - 24) / row_height))
        records = len(self.data)
        pages = math.ceil(records / capacity) if records else 0
        overflow = records % capacity if records else 0
        last_page_records = overflow or (capacity if records else 0)
        unused_rows = capacity - last_page_records if records else 0
        return {"selected": selected, "page_size": page_size, "capacity": capacity, "pages": pages,
                "overflow": overflow, "last_page_records": last_page_records, "unused_rows": unused_rows,
                "records": records, "printable_width": width - int(self.left_margin_var.get()) * mm - int(self.right_margin_var.get()) * mm}

    def _guidance(self, metrics):
        selected = len(metrics["selected"])
        notices = []
        if not selected:
            notices.append("Do: select the fields to include before saving the report.")
        else:
            notices.append(f"Do: keep the {selected} selected fields relevant to this report.")
        maximum = 6 if self.orientation_var.get() == "Portrait" else 10
        if selected > maximum:
            if self.orientation_var.get() == "Portrait":
                notices.append(f"Recommendation: use Landscape because {selected} columns are selected. Portrait is best kept to about {maximum} columns.")
            else:
                notices.append(f"Recommendation: {selected} columns may be crowded in landscape. Use A3 or remove less-important columns for a clearer business report.")
        elif selected:
            notices.append("Good: the current field count should remain readable at the selected page size.")
        if min(int(self.top_margin_var.get()), int(self.bottom_margin_var.get()), int(self.left_margin_var.get()), int(self.right_margin_var.get())) < 5:
            notices.append("Avoid: margins below 5 mm; printers may clip the report edges.")
        else:
            notices.append("Do: keep at least 5 mm margins for reliable printing.")
        if int(self.font_size_body_var.get()) < 8:
            notices.append("Avoid: body text below 8 pt for a report that will be printed.")
        if metrics["pages"] > 1 and 0 < metrics["overflow"] <= 4:
            if int(self.top_margin_var.get()) > 5 or int(self.bottom_margin_var.get()) > 5:
                notices.append(f"Recommendation: the last page has only {metrics['overflow']} record(s). Reduce the top/bottom margins slightly, or reduce body font by 1 pt, to try to avoid a mostly empty final page.")
            else:
                notices.append(f"Note: the last page has only {metrics['overflow']} record(s). Margins are already tight, so consider reducing body font by 1 pt or using a larger page size.")
        elif metrics["pages"] == 1 and metrics["records"] and metrics["unused_rows"] >= metrics["capacity"] * 0.6:
            notices.append("Recommendation: this is a short one-page report. Increase body font slightly or increase vertical margins for a more balanced printed page.")
        if metrics["records"] == 0:
            notices.append("Note: this filtered report contains no records; change the report filter before export.")
        else:
            notices.append(f"Preview: {metrics['records']} records, about {metrics['pages']} page(s), {metrics['capacity']} records per page, and {metrics['overflow']} overflow record(s) on the final page.")
        return "\n\n".join(notices)

    def update_field_count(self, event=None):
        if not hasattr(self, "pdf_win"):
            return
        metrics = self._layout_metrics()
        selected = len(metrics["selected"])
        total = len(self.pdf_field_vars)
        self.lbl_field_count.config(text=f"Selected Fields : {selected} / {total}")
        values = {"name": self.report_title, "records": str(metrics["records"]), "fields": str(selected),
                  "page_size": self.page_size_var.get(), "orientation": self.orientation_var.get(),
                  "capacity": str(metrics["capacity"]), "overflow": f"{metrics['overflow']} record(s)",
                  "pages": str(metrics["pages"]), "status": "Ready" if selected else "Select fields"}
        for key, value in values.items():
            self.summary_labels[key].config(text=value)
        self.guidance_label.config(text=self._guidance(metrics))
        self.show_preview()

    def clear_all_fields(self):
        for variable in self.pdf_field_vars.values():
            variable.set(False)
        self.update_field_count()

    def select_all_fields(self):
        for variable in self.pdf_field_vars.values():
            variable.set(True)
        self.update_field_count()

    def show_preview(self):
        if not hasattr(self, "preview_canvas"):
            return
        metrics = self._layout_metrics()
        canvas = self.preview_canvas
        canvas.delete("all")
        canvas_width, canvas_height = 320, 253
        page_width, page_height = metrics["page_size"]
        scale = min((canvas_width - 24) / page_width, (canvas_height - 20) / page_height)
        preview_width, preview_height = page_width * scale, page_height * scale
        x1, y1 = (canvas_width - preview_width) / 2, (canvas_height - preview_height) / 2
        x2, y2 = x1 + preview_width, y1 + preview_height
        canvas.create_rectangle(x1, y1, x2, y2, outline="black", fill="white")
        left = x1 + int(self.left_margin_var.get()) * mm * scale
        right = x2 - int(self.right_margin_var.get()) * mm * scale
        top = y1 + int(self.top_margin_var.get()) * mm * scale
        bottom = y2 - int(self.bottom_margin_var.get()) * mm * scale
        canvas.create_rectangle(left, top, right, bottom, outline="#cc3333", width=1)
        title = self.company_title_var.get().strip() or self.report_title
        if self.company_heading_var.get():
            company, _logo = self.get_company_details()
            canvas.create_text((left + right) / 2, top + 7, text=company or "Company Name", font=("Arial", 6, "bold"))
        canvas.create_text((left + right) / 2, top + 16, text=title[:35], font=("Arial", 6, "bold"))
        header_y = top + 25
        canvas.create_rectangle(left, header_y, right, header_y + 9, fill="#808080", outline="black")
        selected = metrics["selected"]
        shown = selected[:6 if self.orientation_var.get() == "Portrait" else 9]
        if shown:
            width = (right - left) / len(shown)
            for index, name in enumerate(shown):
                x = left + width * index + width / 2
                canvas.create_text(x, header_y + 4, text=name[:10], fill="white", font=("Arial", 4, "bold"))
                canvas.create_line(left + width * index, header_y, left + width * index, bottom)
            for index in range(5):
                y = header_y + 9 + index * 8
                if y < bottom:
                    canvas.create_line(left, y, right, y, fill="#999999")
        else:
            canvas.create_text((left + right) / 2, (top + bottom) / 2, text="Select report fields", fill="grey", font=("Arial", 8, "bold"))
        self.preview_info_label.config(
            text=(f"Red border = printable area. Orientation: {self.orientation_var.get()} | "
                  f"Records: {metrics['records']} | Per page: {metrics['capacity']} | "
                  f"Estimated pages: {metrics['pages']} | Overflow: {metrics['overflow']} record(s).")
        )

    def _safe_font(self, font_name):
        try:
            FontManager.register_reportlab(font_name)
        except Exception:
            pass
        return font_name if font_name in pdfmetrics.getRegisteredFontNames() else "Helvetica"

    def _footer(self, canvas, document):
        canvas.saveState()
        generated_at = getattr(self, "_active_report_context", {}).get("generated_at", datetime.now())
        if self.show_date_var.get():
            canvas.setFont("Helvetica", 7)
            canvas.drawString(document.leftMargin, 7 * mm, generated_at.strftime("Generated: %d-%b-%Y %H:%M"))
        if self.show_page_var.get():
            canvas.setFont("Helvetica", 7)
            canvas.drawRightString(document.pagesize[0] - document.rightMargin, 7 * mm, f"Page {document.page}")
        canvas.restoreState()

    def generate_pdf(self):
        context = self._build_report_context()
        self._active_report_context = context
        metrics = context["metrics"]
        selected = context["selected"]
        report_rows = context["rows"]
        if not selected:
            messagebox.showwarning("FTMS PRO", "Select one or more report fields before saving the PDF.", parent=self.pdf_win)
            return None
        if not self.data:
            messagebox.showwarning("FTMS PRO", "There are no report records to export.", parent=self.pdf_win)
            return None
        file_path = filedialog.asksaveasfilename(defaultextension=".pdf", filetypes=[("PDF Files", "*.pdf")], title="FTMS PRO: Save PDF Report")
        if not file_path:
            return None
        try:
            company_font = self._safe_font(self.font_comphead_var.get())
            title_font = self._safe_font(self.font_compreport_var.get())
            body_font = self._safe_font(self.font_body_var.get())
            styles = getSampleStyleSheet()
            company_style = ParagraphStyle("FTMSCompany", parent=styles["Normal"], fontName=company_font,
                                           fontSize=int(self.font_size_comphead_var.get()), alignment=TA_CENTER, spaceAfter=5)
            title_style = ParagraphStyle("FTMSTitle", parent=styles["Normal"], fontName=title_font,
                                         fontSize=int(self.font_size_compreport_var.get()), alignment=TA_CENTER, spaceAfter=8)
            body_style = ParagraphStyle("FTMSBody", parent=styles["Normal"], fontName=body_font,
                                        fontSize=int(self.font_size_body_var.get()), leading=max(9, int(self.font_size_body_var.get()) + 2))
            header_style = ParagraphStyle("FTMSHeader", parent=body_style, fontSize=body_style.fontSize + 2)
            document = SimpleDocTemplate(file_path, pagesize=metrics["page_size"],
                                         topMargin=int(self.top_margin_var.get()) * mm, bottomMargin=int(self.bottom_margin_var.get()) * mm,
                                         leftMargin=int(self.left_margin_var.get()) * mm, rightMargin=int(self.right_margin_var.get()) * mm)
            story = []
            company_name = context["company_name"]
            logo_path = context["logo_path"]
            if self.company_logo_var.get() and logo_path and os.path.exists(logo_path):
                image = ReportLabImage(logo_path)
                image._restrictSize(45 * mm, 20 * mm)
                image.hAlign = "CENTER"
                story.append(image)
            if self.company_heading_var.get() and company_name:
                story.append(Paragraph(html.escape(company_name), company_style))
            title = context["title"]
            story.append(Paragraph(html.escape(title), title_style))
            table_data = [[Paragraph(f"<b>{html.escape(str(column).title())}</b>", header_style) for column in selected]]
            for row in report_rows:
                table_data.append([Paragraph(html.escape(str(value)), body_style) for value in row])
            table = Table(table_data, colWidths=[metrics["printable_width"] / len(selected)] * len(selected), repeatRows=1, hAlign="CENTER")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f2f2f2")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.black), ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("GRID", (0, 0), (-1, -1), 0.35, colors.black),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.whitesmoke, colors.HexColor("#eaf2f8")]),
                ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.extend([Spacer(1, 2), table])
            document.build(story, onFirstPage=self._footer, onLaterPages=self._footer)
            messagebox.showinfo("FTMS PRO", "PDF report exported successfully.", parent=self.pdf_win)
            return file_path
        except Exception as error:
            messagebox.showerror("FTMS PRO: PDF Error", str(error), parent=self.pdf_win)
            return None

    def generate_excel(self):
        context = self._build_report_context()
        selected = context["selected"]
        report_rows = context["rows"]
        if not selected:
            messagebox.showwarning("FTMS PRO", "Select one or more report fields before saving the Excel report.", parent=self.pdf_win)
            return None
        if not report_rows:
            messagebox.showwarning("FTMS PRO", "There are no report records to export.", parent=self.pdf_win)
            return None
        file_path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel Files", "*.xlsx")], title="FTMS PRO: Save Excel Report")
        if not file_path:
            return None
        try:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Vehicle Report"
            sheet.sheet_view.showGridLines = False
            max_column = len(selected)
            last_column = get_column_letter(max_column)
            row_number = 1

            if self.company_logo_var.get() and context["logo_path"] and os.path.exists(context["logo_path"]):
                try:
                    logo = ExcelImage(context["logo_path"])
                    logo.width, logo.height = 170, 75
                    center_column = max(1, (max_column + 1) // 2)
                    logo_column = get_column_letter(center_column)
                    sheet.add_image(logo, f"{logo_column}1")
                    row_number = 6
                except Exception:
                    pass
            if self.company_heading_var.get() and context["company_name"]:
                sheet.merge_cells(f"A{row_number}:{last_column}{row_number}")
                company_cell = sheet.cell(row=row_number, column=1, value=context["company_name"])
                company_cell.font = Font(name=context["company_font"], size=int(self.font_size_comphead_var.get()), bold=True)
                company_cell.alignment = Alignment(horizontal="center")
                row_number += 1

            sheet.merge_cells(f"A{row_number}:{last_column}{row_number}")
            title_cell = sheet.cell(row=row_number, column=1, value=context["title"])
            title_cell.font = Font(name=context["title_font"], size=int(self.font_size_compreport_var.get()), bold=True)
            title_cell.alignment = Alignment(horizontal="center")
            row_number += 1

            metadata = context["metrics"]
            sheet.merge_cells(f"A{row_number}:{last_column}{row_number}")
            metadata_cell = sheet.cell(
                row=row_number,
                column=1,
                value=(f"Generated: {context['generated_at'].strftime('%d-%b-%Y %H:%M')} | "
                       f"Orientation: {self.orientation_var.get()} | Records: {metadata['records']} | "
                       f"Records/Page: {metadata['capacity']} | Estimated Pages: {metadata['pages']}")
            )
            metadata_cell.font = Font(name=context["body_font"], size=max(8, int(self.font_size_body_var.get()) - 1), italic=True)
            metadata_cell.alignment = Alignment(horizontal="center")
            row_number += 2

            header_fill = PatternFill("solid", fgColor="F2F2F2")
            thin_border = Border(
                left=Side(style="thin", color="000000"), right=Side(style="thin", color="000000"),
                top=Side(style="thin", color="000000"), bottom=Side(style="thin", color="000000"),
            )
            for column_number, heading in enumerate(selected, start=1):
                cell = sheet.cell(row=row_number, column=column_number, value=str(heading).title())
                cell.font = Font(name=context["body_font"], size=int(self.font_size_body_var.get()) + 2, bold=True, color="000000")
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = thin_border
            header_row = row_number

            for row in report_rows:
                row_number += 1
                for column_number, value in enumerate(row, start=1):
                    cell = sheet.cell(row=row_number, column=column_number, value=value)
                    cell.font = Font(name=context["body_font"], size=int(self.font_size_body_var.get()))
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    cell.border = thin_border

            for column_number, heading in enumerate(selected, start=1):
                values = [str(heading)] + [str(row[column_number - 1]) for row in report_rows]
                sheet.column_dimensions[get_column_letter(column_number)].width = min(32, max(12, max(map(len, values)) + 2))
            sheet.freeze_panes = f"A{header_row + 1}"
            sheet.auto_filter.ref = f"A{header_row}:{last_column}{row_number}"
            sheet.page_setup.orientation = self.orientation_var.get().lower()
            sheet.page_margins.left = int(self.left_margin_var.get()) / 25.4
            sheet.page_margins.right = int(self.right_margin_var.get()) / 25.4
            sheet.page_margins.top = int(self.top_margin_var.get()) / 25.4
            sheet.page_margins.bottom = int(self.bottom_margin_var.get()) / 25.4
            if self.show_date_var.get():
                sheet.oddFooter.left.text = f"Generated: {context['generated_at'].strftime('%d-%b-%Y %H:%M')}"
            if self.show_page_var.get():
                sheet.oddFooter.right.text = "Page &P"
            workbook.save(file_path)
            messagebox.showinfo("FTMS PRO", "Excel report exported successfully.", parent=self.pdf_win)
            return file_path
        except Exception as error:
            messagebox.showerror("FTMS PRO: Excel Error", str(error), parent=self.pdf_win)
            return None

    def print_pdf(self):
        file_path = self.generate_pdf()
        if file_path:
            try:
                os.startfile(file_path, "print")
            except OSError as error:
                messagebox.showerror("FTMS PRO: Print Error", str(error), parent=self.pdf_win)

    def export_document_pdf(self):
        messagebox.showinfo("FTMS PRO", "Document scanning and conversion are handled by GraceScan. This screen exports the current vehicle report.", parent=self.pdf_win)

    def clear_logo_preview(self):
        if hasattr(self, "logo_canvas"):
            self.logo_canvas.delete("all")
            self.logo_canvas.create_text(150, 72, text="No Logo Used", fill="grey", font=("Arial", 16, "bold"))

    def show_company_logo(self):
        _company_name, logo_path = self.get_company_details()
        self.logo_canvas.delete("all")
        if logo_path and os.path.exists(logo_path):
            try:
                image = Image.open(logo_path)
                image.thumbnail((250, 115))
                self.logo_img = ImageTk.PhotoImage(image)
                self.logo_canvas.create_image(150, 72, image=self.logo_img)
                return
            except Exception:
                pass
        self.logo_canvas.create_text(150, 72, text="No Logo Available", fill="grey", font=("Arial", 16, "bold"))

    def toggle_company_logo(self):
        if self.company_logo_var.get():
            self.show_company_logo()
        else:
            self.clear_logo_preview()
        self.update_field_count()

    def pdf_win_close(self):
        if messagebox.askyesno("FTMS PRO", "Are you sure you want to close?", parent=self.pdf_win):
            self.pdf_win.destroy()

    def get_page_size_mm(self):
        width, height = self.PAGE_SIZES[self.page_size_var.get()]
        return round(width / mm, 1), round(height / mm, 1)
